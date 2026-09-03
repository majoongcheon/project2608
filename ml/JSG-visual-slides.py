# -*- coding: utf-8 -*-
"""발표용 시각화 자료 — 그래프를 하나씩 따로 그려 pptx 로 묶는다.

기존 docs/assets/*.png 은 그래프와 해설 문구·화살표·제목이 한 이미지에 함께 구워져
있어 슬라이드에서 글자만 고칠 수 없었다. 이 스크립트는 **그래프 본체와 그에 딸린
축 이름·눈금·범례·값 라벨만** 이미지로 남기고, 제목과 설명은 pptx 의 편집 가능한
텍스트 상자로 분리한다.

  python3 ml/JSG-visual-slides.py

산출: docs/assets/발표용/*.png · docs/발표용_시각화_자료.pptx
읽기 전용 — DB·models·서비스를 건드리지 않는다.
"""
import json, sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'ml' / '.pylibs'))

import numpy as np
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager

font_manager.fontManager.addfont('/System/Library/Fonts/Supplemental/AppleGothic.ttf')
matplotlib.rcParams['font.family'] = 'AppleGothic'
matplotlib.rcParams['axes.unicode_minus'] = False

D = json.loads((ROOT / 'docs/assets/chart_data.json').read_text(encoding='utf-8'))
OUT = ROOT / 'docs/assets/발표용'; OUT.mkdir(parents=True, exist_ok=True)

PAPER, INK, BODY, MUTED = '#FDF9F4', '#2B2724', '#4A443E', '#8A8176'
LINE = '#E4DACC'
CAT1, CAT2, CAT3 = '#eb6834', '#2a78d6', '#1baf7a'
GRAY = '#9A9187'
BURDEN = {1: '#7A2E10', 2: '#B04A1E', 3: '#DC7A45', 4: '#F0A87C', 5: '#F6C9AC'}
NAME = {1: '최고부담군', 2: '고부담군', 3: '중간부담군', 4: '저부담군', 5: '부담 없음'}

SLIDES = []   # (png 파일명, 제목, 설명 한 줄)


def fig(w, h, grid='y'):
    f, ax = plt.subplots(figsize=(w, h))
    f.patch.set_facecolor(PAPER); ax.set_facecolor(PAPER)
    for s in ('top', 'right'): ax.spines[s].set_visible(False)
    for s in ('left', 'bottom'): ax.spines[s].set_color(LINE)
    ax.tick_params(colors=MUTED, length=0, labelsize=11)
    ax.set_axisbelow(True)
    if grid: ax.grid(axis=grid, color=LINE, lw=0.8)
    return f, ax


def save(f, name, title, note):
    f.tight_layout()
    f.savefig(OUT / name, dpi=200, facecolor=PAPER); plt.close(f)
    SLIDES.append((name, title, note))
    print('  ', name)


# ── 01. 돌봄부담 5구간 분포 ────────────────────────────────────────
CNT = {1: 464, 2: 1165, 3: 916, 4: 378, 5: 77}
tot = sum(CNT.values())
f, ax = fig(11, 2.6, grid=None)
left = 0
for lv in [1, 2, 3, 4, 5]:
    pct = CNT[lv] / tot * 100
    ax.barh([0], [pct], left=left, height=0.55, color=BURDEN[lv], edgecolor=PAPER, lw=2,
            label=f'{NAME[lv]}  {CNT[lv]:,}가구 ({pct:.1f}%)')
    if pct > 6:
        ax.text(left + pct / 2, 0, f'{pct:.1f}%', ha='center', va='center',
                fontsize=12, color='white' if lv <= 2 else INK, fontweight='bold')
    left += pct
ax.set_xlim(0, 100); ax.set_ylim(-0.5, 0.45)
ax.set_yticks([]); ax.set_xticks([0, 25, 50, 75, 100])
ax.set_xticklabels(['0', '25', '50', '75', '100'])
ax.set_xlabel('구성비 (%)', fontsize=12, color=BODY)
ax.legend(frameon=False, fontsize=10.5, ncol=3, loc='upper center',
          bbox_to_anchor=(0.5, -0.42), handlelength=1.2)
save(f, 'V01-부담분포.png', '돌봄부담 5구간 분포',
     '2024 발달장애인 일과 삶 실태조사 3,000가구 · 내부 라벨은 1이 최고부담인 역방향 척도')

# ── 02. 5구간 건수 세로 막대 (Summary) ────────────────────────────
f, ax = fig(8.4, 4.4)
xs = [NAME[v] for v in [1, 2, 3, 4, 5]]
vs = [CNT[v] for v in [1, 2, 3, 4, 5]]
b = ax.bar(xs, vs, width=0.6, color=[BURDEN[v] for v in [1, 2, 3, 4, 5]])
for r, v in zip(b, vs):
    ax.text(r.get_x() + r.get_width() / 2, v + 25, f'{v:,}', ha='center',
            fontsize=12, fontweight='bold', color=INK)
ax.set_ylabel('가구 수', fontsize=12, color=BODY)
ax.set_ylim(0, max(vs) * 1.16)
ax.tick_params(axis='x', labelsize=11.5, colors=BODY)
save(f, 'V02-부담분포-건수.png', '돌봄부담 5구간 가구 수',
     '전체 3,000가구 · 5구간(부담 없음)은 77가구뿐이라 층화 분할이 필요했다')

# ── 03. 단독 설명력 순위 ──────────────────────────────────────────
LAB = {'help_needed_hours': '도움 필요 시간', 'understands_work_meaning': '근로 의미 이해도',
       'can_work_standard_job': '통상근로 가능', 'severe_dd_job_willingness': '최중증 취업의사',
       'family_support_for_employment': '가족 취업지지', 'daily_routine_satisfaction': '일과 만족도',
       'overall_health': '전반적 건강', 'wants_person_employed': '취업 희망',
       'past_employment_exp': '과거 취업경험', 'employment_status': '종사상 지위',
       'job_ability_mobility': '이동 능력', 'is_employed': '취업 여부',
       'job_ability_strength': '근력·체력', 'school_helpfulness': '학교교육 도움',
       'household_head_type': '가구주 유형', 'relation_to_person': '당사자와 관계',
       'last_job_quit_reason': '퇴사 이유', 'caregiver_age': '보호자 나이',
       'age_disability_suspected': '장애의심 나이', 'caregiver_life_satisfaction': '보호자 삶 만족도',
       'care_difficulty_top1': '돌봄 어려움 1순위', 'needed_care_service_type': '필요 돌봄서비스',
       'work_care_gap_hours': '월평균 돌봄공백', 'work_care_gap_exp': '근로중 돌봄공백',
       'integrated_care_awareness': '통합돌봄 인지'}
top = D['mi'][:18]
f, ax = fig(10.4, 6.0, grid='x')
ys = np.arange(len(top))[::-1]
sel7 = set(D['selected7'])
for yy, r in zip(ys, top):
    if r['excluded']: c, hatch = GRAY, '///'
    elif r['feature'] in sel7: c, hatch = CAT1, None
    else: c, hatch = '#D8CFC3', None
    ax.barh([yy], [r['mi']], height=0.7, color=c, hatch=hatch, edgecolor=PAPER)
    ax.text(r['mi'] + 0.25, yy, f"{r['mi']:.2f}", va='center', fontsize=9.5, color=MUTED)
ax.set_yticks(ys)
ax.set_yticklabels([LAB.get(r['feature'], r['feature']) for r in top], fontsize=10.5, color=BODY)
ax.set_xlabel('단독 설명력  MI / H  (%)', fontsize=12, color=BODY)
ax.set_xlim(0, 21)
h = [plt.Rectangle((0, 0), 1, 1, color=CAT1),
     plt.Rectangle((0, 0), 1, 1, color='#D8CFC3'),
     plt.Rectangle((0, 0), 1, 1, color=GRAY, hatch='///')]
ax.legend(h, ['채택 7문항', '잔존(미채택)', 'Data Leakage 배제 6개'],
          frameon=False, fontsize=10.5, loc='lower right')
save(f, 'V03-설명력순위.png', '변수별 단독 설명력과 누수 배제',
     'train 2,398건 기준 상위 18개 · 배제 1위(보호자 삶 만족도 19.06%)는 남은 최고치의 2.3배')

# ── 04. PCA 산점도 — 부담 구간별 작은 다중 패널 ────────────────────
P = D['pca']; lab = np.array(P['label']); evr = P['evr']
px, py = np.array(P['x']), np.array(P['y'])
f = plt.figure(figsize=(11.6, 2.9)); f.patch.set_facecolor(PAPER)
for c, lv in enumerate([1, 2, 3, 4, 5]):
    ax = f.add_subplot(1, 5, c + 1); ax.set_facecolor(PAPER)
    m = lab == lv
    ax.scatter(px[~m], py[~m], s=3.5, color='#DFD8CE', alpha=0.6, linewidths=0)
    ax.scatter(px[m], py[m], s=5.5, color=CAT1, alpha=0.85, linewidths=0)
    ax.set_xticks([]); ax.set_yticks([]); ax.set_aspect('equal')
    for sp in ax.spines.values(): sp.set_color(LINE)
    ax.set_title(f'{NAME[lv]}  n={m.sum():,}', fontsize=10.5, color=INK, pad=5)
    if c == 0:
        ax.set_ylabel(f'PC2 ({evr[1]*100:.1f}%)', fontsize=10, color=BODY)
    ax.set_xlabel(f'PC1 ({evr[0]*100:.1f}%)', fontsize=10, color=BODY)
f.tight_layout()
f.savefig(OUT / 'V04-PCA-구간별산점도.png', dpi=200, facecolor=PAPER); plt.close(f)
SLIDES.append(('V04-PCA-구간별산점도.png', '주성분 평면에서 본 부담 구간',
               '회색 = 전체 2,398건, 주황 = 해당 구간 · 어느 구간도 특정 영역에 몰리지 않는다'))
print('   V04-PCA-구간별산점도.png')

# ── 05. 실루엣 계수 ───────────────────────────────────────────────
C = D['cluster']
f, ax = fig(8.4, 4.4)
ax.plot(C['k'], C['silhouette'], color=CAT2, lw=2.4, marker='o', ms=6, label='실루엣 계수')
ax.axhline(0.25, color=GRAY, lw=1.5, ls=(0, (4, 3)), label='구조 판정 하한 0.25')
ax.scatter([2], [C['silhouette'][0]], s=170, facecolor=PAPER, edgecolor=CAT2, lw=2.6, zorder=5)
for k, v in zip(C['k'], C['silhouette']):
    ax.text(k, v + 0.011, f'{v:.3f}', ha='center', fontsize=9, color=MUTED)
ax.set_xlabel('군집 수  k', fontsize=12, color=BODY)
ax.set_ylabel('실루엣 계수', fontsize=12, color=BODY)
ax.set_ylim(0, 0.30); ax.set_xticks(C['k'])
ax.legend(frameon=False, fontsize=10.5, loc='upper right')
save(f, 'V05-실루엣계수.png', '군집 수별 실루엣 계수',
     'y 를 빼고 38변수만으로 KMeans · 전 구간 0.25 미만 = 뚜렷한 덩어리 구조 없음 · 최고는 k=2')

# ── 06. ARI ──────────────────────────────────────────────────────
f, ax = fig(8.4, 4.4)
ax.plot(C['k'], C['ari'], color=CAT3, lw=2.4, marker='s', ms=6, label='ARI (군집 ↔ 부담 구간 일치도)')
for k, v in zip(C['k'], C['ari']):
    ax.text(k, v + 0.0016, f'{v:.3f}', ha='center', fontsize=9, color=MUTED)
ax.axhline(0, color=GRAY, lw=1.2)
ax.set_xlabel('군집 수  k', fontsize=12, color=BODY)
ax.set_ylabel('조정 랜드 지수 (ARI)', fontsize=12, color=BODY)
ax.set_ylim(-0.005, 0.045); ax.set_xticks(C['k'])
ax.legend(frameon=False, fontsize=10.5, loc='upper right')
save(f, 'V06-ARI.png', '군집과 실제 부담 구간의 일치도',
     '0 = 무작위와 다를 바 없음 · 전 구간 0.03 이하로 군집 축은 부담이 아니다')

# ── 07·08. 문항 수 곡선 ───────────────────────────────────────────
cv = sorted(D['curve'], key=lambda r: r['k'])
ks = [r['k'] for r in cv]; f1 = [r['f1'] for r in cv]; rc = [r['recall'] for r in cv]
f, ax = fig(10.4, 4.4)
ax.plot(ks, f1, color=CAT1, lw=2.4, marker='o', ms=4.5, label='macro F1')
ax.axvline(7, color=INK, lw=1.3, ls=(0, (2, 2)), label='채택 지점 k=7')
ax.scatter([7], [0.3383], s=190, facecolor=PAPER, edgecolor=CAT1, lw=2.8, zorder=6)
ax.set_xlabel('문항 수  k', fontsize=12, color=BODY)
ax.set_ylabel('macro F1', fontsize=12, color=BODY)
ax.set_ylim(0.28, 0.355); ax.invert_xaxis()
ax.legend(frameon=False, fontsize=10.5, loc='lower left')
save(f, 'V07-문항수-F1.png', '후진 제거 — 문항 수별 macro F1',
     '38개에서 하나씩 제거하며 5-fold 측정 · k=7 에서 최고 0.3383, k=6 에서 급락')

f, ax = fig(10.4, 4.4)
ax.plot(ks, rc, color=CAT2, lw=2.4, marker='s', ms=4.5, label='고부담 재현율')
ax.axhline(0.70, color=GRAY, lw=1.5, ls=(0, (4, 3)), label='정지 조건 하한 0.70')
ax.axvline(7, color=INK, lw=1.3, ls=(0, (2, 2)), label='채택 지점 k=7')
ax.scatter([7], [0.7843], s=190, facecolor=PAPER, edgecolor=CAT2, lw=2.8, zorder=6)
ax.set_xlabel('문항 수  k', fontsize=12, color=BODY)
ax.set_ylabel('고부담 재현율', fontsize=12, color=BODY)
ax.set_ylim(0.68, 0.82); ax.invert_xaxis()
ax.legend(frameon=False, fontsize=10.5, loc='lower left')
save(f, 'V08-문항수-재현율.png', '후진 제거 — 문항 수별 고부담 재현율',
     '하한 0.70 을 한 번도 깨지 않았다 · k=7 에서 0.7843')

# ── 09~11. 혼동행렬 2×2 ───────────────────────────────────────────
CF = D['confusion']
for k in ['38', '7', '6']:
    cm = np.array(CF[k]['cm2'], dtype=float)
    pct = cm / cm.sum(axis=1, keepdims=True) * 100
    f, ax = plt.subplots(figsize=(5.0, 4.6))
    f.patch.set_facecolor(PAPER); ax.set_facecolor(PAPER)
    ax.imshow(pct, cmap='OrRd', vmin=0, vmax=100, alpha=0.85)
    for i in range(2):
        for j in range(2):
            ax.text(j, i, f'{int(cm[i,j])}\n{pct[i,j]:.0f}%', ha='center', va='center',
                    fontsize=15, fontweight='bold',
                    color='white' if pct[i, j] > 55 else INK)
    ax.set_xticks([0, 1]); ax.set_yticks([0, 1])
    ax.set_xticklabels(['고부담', '아님'], fontsize=12, color=BODY)
    ax.set_yticklabels(['고부담', '아님'], fontsize=12, color=BODY)
    ax.set_xlabel('예측', fontsize=12.5, color=BODY)
    ax.set_ylabel('실제', fontsize=12.5, color=BODY)
    ax.tick_params(length=0)
    for s in ax.spines.values(): s.set_visible(False)
    save(f, f'V09-혼동행렬2x2-{k}문항.png', f'고부담 여부 혼동행렬 — {k}문항',
         f'교차검증 2,398건 · 재현율 {CF[k]["recall"]:.4f} · macro F1 {CF[k]["macro_f1"]:.4f} · 셀 안은 건수와 행 기준 비율')

# ── 12~13. 혼동행렬 5×5 ───────────────────────────────────────────
n5 = [NAME[v] for v in [1, 2, 3, 4, 5]]
for k in ['7', '6']:
    cm5 = np.array(CF[k]['cm5'], dtype=float)
    row = cm5 / cm5.sum(axis=1, keepdims=True) * 100
    f, ax = plt.subplots(figsize=(6.4, 5.6))
    f.patch.set_facecolor(PAPER); ax.set_facecolor(PAPER)
    ax.imshow(row, cmap='OrRd', vmin=0, vmax=75, alpha=0.85)
    for i in range(5):
        for j in range(5):
            ax.text(j, i, f'{int(cm5[i,j])}', ha='center', va='center', fontsize=12,
                    fontweight='bold' if i == j else 'normal',
                    color='white' if row[i, j] > 48 else INK)
    ax.set_xticks(range(5)); ax.set_yticks(range(5))
    ax.set_xticklabels(n5, fontsize=10, color=BODY, rotation=30, ha='right')
    ax.set_yticklabels(n5, fontsize=10, color=BODY)
    ax.set_xlabel('예측', fontsize=12.5, color=BODY)
    ax.set_ylabel('실제', fontsize=12.5, color=BODY)
    ax.tick_params(length=0)
    for sp in ax.spines.values(): sp.set_visible(False)
    rec1 = cm5[0, 0] / cm5[0].sum()
    pred1 = int(cm5[:, 0].sum())
    save(f, f'V12-혼동행렬5x5-{k}문항.png', f'5구간 혼동행렬 — {k}문항',
         f'교차검증 2,398건 · 최고부담군 재현율 {rec1:.3f} · 최고부담군으로 예측한 건수 {pred1}건 · macro F1 {CF[k]["macro_f1"]:.4f}')

# ── 14~17. 7문항 vs 14문항 지표 4종 ───────────────────────────────
METRICS = [
    ('V14-7vs14-F1.png',      'macro F1 (test 602건)', 0.3416, 0.3420, '{:.4f}', 'macro F1',
     '차이 0.0004 — 사실상 같다'),
    ('V15-7vs14-재현율.png',   '고부담 재현율 (test 602건)', 0.7730, 0.7669, '{:.4f}', '고부담 재현율',
     '문항을 늘렸는데 7문항이 오히려 높다'),
    ('V16-7vs14-판정불가.png', '판정 불가 비율 (test 602건)', 2.49, 4.20, '{:.2f}%', '판정 불가 비율 (%)',
     '낮을수록 좋다 — 14문항이 1.7배 · 원인은 78범주짜리 나이 변수'),
    ('V17-7vs14-열수.png',     '원-핫 설계행렬 열 수', 53, 224, '{:.0f}', '열 수',
     '낮을수록 단순하다 — 4.2배 차이'),
]
for fn, title, v7, v14, fmt, ylab, note in METRICS:
    f, ax = fig(5.2, 4.4)
    b = ax.bar(['7문항', '14문항'], [v7, v14], width=0.5, color=[CAT1, CAT2])
    top_v = max(v7, v14)
    for r, v in zip(b, [v7, v14]):
        ax.text(r.get_x() + r.get_width() / 2, v + top_v * 0.035, fmt.format(v),
                ha='center', fontsize=14, fontweight='bold', color=INK)
    ax.set_ylim(0, top_v * 1.25)
    ax.set_ylabel(ylab, fontsize=12, color=BODY)
    ax.tick_params(axis='x', labelsize=12.5, colors=BODY)
    save(f, fn, title, note)

# ── pptx 로 묶는다 ────────────────────────────────────────────────
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image

W, H = Inches(13.333), Inches(7.5)
prs = Presentation(); prs.slide_width, prs.slide_height = W, H
blank = prs.slide_layouts[6]


def rgb(h): return RGBColor.from_string(h.lstrip('#').upper())


def textbox(slide, l, t, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = rgb(color); r.font.name = 'AppleGothic'
    return tb


# 표지
s = prs.slides.add_slide(blank)
s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(PAPER)
textbox(s, Inches(0.9), Inches(2.5), Inches(11.5), Inches(1.2),
        '발표용 시각화 자료', 40, INK, True)
textbox(s, Inches(0.9), Inches(3.6), Inches(11.5), Inches(1.6),
        '곁 — 돌봄의 무게를 읽다  ·  그래프 한 장에 하나씩\n'
        '각 장은 그래프 본체와 그에 딸린 축·눈금·범례만 이미지이고, 제목과 설명은 편집 가능한 텍스트 상자다.',
        15, BODY)
textbox(s, Inches(0.9), Inches(6.4), Inches(11.5), Inches(0.5),
        'Pioneer 3 팀 · 2026-09-03 · 생성 스크립트 ml/JSG-visual-slides.py', 12, MUTED)

for i, (png, title, note) in enumerate(SLIDES, 1):
    s = prs.slides.add_slide(blank)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(PAPER)
    textbox(s, Inches(0.7), Inches(0.42), Inches(11.9), Inches(0.7), title, 26, INK, True)
    textbox(s, Inches(0.7), Inches(1.12), Inches(11.9), Inches(0.5), note, 12.5, MUTED)
    textbox(s, Inches(11.9), Inches(6.9), Inches(1.0), Inches(0.35),
            f'{i:02d}', 11, MUTED, align=PP_ALIGN.RIGHT)

    iw, ih = Image.open(OUT / png).size
    box_l, box_t = Inches(0.7), Inches(1.75)
    box_w, box_h = Inches(11.9), Inches(5.0)
    sc = min(box_w / iw, box_h / ih)
    w, h = int(iw * sc), int(ih * sc)
    s.shapes.add_picture(str(OUT / png), box_l + int((box_w - w) / 2),
                         box_t + int((box_h - h) / 2), width=w, height=h)

dst = ROOT / 'docs/발표용_시각화_자료.pptx'
prs.save(dst)
print(f'\n슬라이드 {len(prs.slides.__iter__.__self__._sldIdLst)}장 → {dst}')
